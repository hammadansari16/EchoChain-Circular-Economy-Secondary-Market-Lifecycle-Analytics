"""Export EchoChain results to a formatted .xlsx workbook, Power-BI-import-ready."""
import os
import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from paths import OUTPUT_DIR, DB_PATH

conn = sqlite3.connect(DB_PATH)
component_scores = pd.read_sql("SELECT * FROM component_scores", conn)
model_scores = pd.read_sql("SELECT * FROM model_scores", conn)
matched = pd.read_sql("SELECT * FROM matched_listings", conn)
conn.close()

wb = Workbook()
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")


def write_sheet(ws, df, title):
    ws.title = title
    for c, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c, value=col.replace("_", " ").title())
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(df.itertuples(index=False), start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val).font = Font(name=FONT)
    for c, col in enumerate(df.columns, start=1):
        width = max(12, min(40, df[col].astype(str).map(len).max() + 2, len(col) + 4))
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


ws1 = wb.active
write_sheet(ws1, model_scores, "Model Circularity Scores")

ws2 = wb.create_sheet()
write_sheet(ws2, component_scores[[
    "model", "component", "unit_cost_usd", "critical_component",
    "warranty_failure_rate_pct", "median_resale_price_usd",
    "resale_retention_pct", "listing_count",
]], "Component Detail")

ws3 = wb.create_sheet()
write_sheet(ws3, matched[[
    "listing_id", "platform", "matched_model", "matched_component",
    "price_usd", "condition", "match_score",
]], "Matched Listings")

out_path = os.path.join(OUTPUT_DIR, "EchoChain_Circularity_Report.xlsx")
wb.save(out_path)
print("Saved:", out_path)
